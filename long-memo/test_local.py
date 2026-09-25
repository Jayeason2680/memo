"""Offline checks. No OpenAI key or network call is used."""
import os
import hashlib
import tempfile
import unittest
import zipfile
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
import httpx

_data = tempfile.TemporaryDirectory()
os.environ["MEMO_DATA_DIR"] = _data.name
os.environ["MEMO_PASSWORD"] = "test-password-long-enough"
os.environ["MEMO_SESSION_SECRET"] = "test-secret-at-least-thirty-two-characters-long"

from fastapi.testclient import TestClient
import app as service
import pipeline


class MemoBetaChecks(unittest.TestCase):
    def create_upload(self, size):
        response = self.client.post("/api/jobs", json={"title": "Upload audit", "filename": "audit.m4a", "size": size, "mode": "meeting"}, headers=self.headers)
        self.assertEqual(response.status_code, 200)
        return response.json()["id"]

    def test_210_mib_upload_preserves_every_byte(self):
        total = 210 * 1024 * 1024
        job_id = self.create_upload(total)
        expected = hashlib.sha256()
        offset = 0
        while offset < total:
            payload = bytes([(offset // service.CHUNK_LIMIT) % 251]) * min(service.CHUNK_LIMIT, total - offset)
            expected.update(payload)
            response = self.client.put(f"/api/jobs/{job_id}/upload", content=payload,
                                       headers={**self.headers, "X-Upload-Offset": str(offset)})
            self.assertEqual(response.status_code, 200)
            offset += len(payload)
            self.assertEqual(response.json()["offset"], offset)
        with (pipeline.folder(job_id) / "original").open("rb") as stream:
            self.assertEqual(hashlib.file_digest(stream, "sha256").hexdigest(), expected.hexdigest())
        with patch.object(service, "dispatch") as dispatch:
            self.assertEqual(self.client.post(f"/api/jobs/{job_id}/finish", headers=self.headers).status_code, 200)
            dispatch.assert_called_once_with(job_id)

    def test_resume_rejects_missing_bytes_and_recovers_unacknowledged_tail(self):
        job_id = self.create_upload(8)
        url = f"/api/jobs/{job_id}/upload"
        headers = {**self.headers, "X-Upload-Offset": "0"}
        self.assertEqual(self.client.put(url, content=b"abcd", headers=headers).status_code, 200)
        original = pipeline.folder(job_id) / "original"
        original.write_bytes(b"ab")
        headers["X-Upload-Offset"] = "4"
        self.assertEqual(self.client.put(url, content=b"efgh", headers=headers).status_code, 409)
        self.assertEqual(original.read_bytes(), b"ab")
        original.write_bytes(b"abcdSTALE")
        self.assertEqual(self.client.put(url, content=b"efgh", headers=headers).status_code, 200)
        self.assertEqual(original.read_bytes(), b"abcdefgh")
        original.write_bytes(b"abc")
        with patch.object(service, "dispatch") as dispatch:
            self.assertEqual(self.client.post(f"/api/jobs/{job_id}/finish", headers=self.headers).status_code, 409)
            dispatch.assert_not_called()

    def setUp(self):
        self.client = TestClient(service.app, base_url="https://testserver")
        login = self.client.post("/api/login", json={"password": os.environ["MEMO_PASSWORD"]})
        self.assertEqual(login.status_code, 200)
        self.headers = {"X-Memo-CSRF": login.json()["csrf"]}
        self.assertEqual(self.client.get("/health").status_code, 200)

    def tearDown(self):
        self.client.close()

    def test_resumable_upload_and_downloads(self):
        payload = {"title": "Bilingual meeting", "filename": "note.m4a", "size": 8, "mode": "meeting", "terms": "银行"}
        self.assertEqual(self.client.post("/api/jobs", json=payload).status_code, 403)
        created = self.client.post("/api/jobs", json=payload, headers=self.headers)
        self.assertEqual(created.status_code, 200)
        job_id = created.json()["id"]
        url = f"/api/jobs/{job_id}/upload"
        self.assertEqual(self.client.put(url, content=b"abcd", headers={**self.headers, "X-Upload-Offset": "1"}).status_code, 409)
        self.assertEqual(self.client.put(url, content=b"abcd", headers={**self.headers, "X-Upload-Offset": "0"}).json()["offset"], 4)
        self.assertEqual(self.client.get(f"/api/jobs/{job_id}").json()["offset"], 4)
        self.assertEqual(self.client.put(url, content=b"efgh", headers={**self.headers, "X-Upload-Offset": "4"}).json()["offset"], 8)
        with patch.object(service, "dispatch"):
            self.assertEqual(self.client.post(f"/api/jobs/{job_id}/finish", headers=self.headers).status_code, 200)
        job = pipeline.read_job(job_id)
        self.assertEqual((pipeline.folder(job_id) / "original").read_bytes(), b"abcdefgh")
        job["status"] = "ready"
        job["transcribed"] = [{"start": 0, "segments": [
            {"start": 2, "end": 4, "speaker": "Part 1 · speaker_0", "text": "Hello, 你好。"}]}]
        job["report"] = {key: "Test paragraph.\n\nSecond paragraph." for key in pipeline.report_schema()["required"]}
        pipeline.save_job(job)
        pipeline.write_exports(job)
        response = self.client.post(f"/api/jobs/{job_id}/speakers", json={"names": {"Part 1 · speaker_0": "Jayson"}}, headers=self.headers)
        self.assertEqual(response.status_code, 200)
        self.assertIn("Jayson: Hello, 你好。", self.client.get(f"/api/jobs/{job_id}/downloads/transcript.txt").text)
        self.assertEqual(pipeline.read_job(job_id)["report"]["brief_en"], "Test paragraph.\n\nSecond paragraph.")
        archive = self.client.get(f"/api/jobs/{job_id}/downloads/memo-downloads.zip").content
        with zipfile.ZipFile(BytesIO(archive)) as bundle:
            self.assertEqual(set(bundle.namelist()), {"transcript.txt", "comprehensive.docx", "brief.docx"})
        self.assertEqual(self.client.delete(f"/api/jobs/{job_id}", headers=self.headers).status_code, 200)

    def test_speaker_ids_remain_part_specific(self):
        job_id = "a" * 32
        path = pipeline.folder(job_id)
        path.mkdir(exist_ok=True)
        (path / "part-0001.m4a").write_bytes(b"mock")
        job = {"id": job_id, "mode": "meeting", "speaker_refs": []}
        with patch.object(pipeline, "_post", return_value={"segments": [
            {"speaker": "speaker_0", "start": 3, "end": 5, "text": "Confirmed."}]}):
            result = pipeline.transcribe_part(job, {"file": "part-0001.m4a", "index": 1, "start": 600})
        self.assertEqual(result["segments"][0]["speaker"], "Part 2 · speaker_0")
        self.assertEqual(result["segments"][0]["start"], 603)

    def test_audio_splits_at_nearby_silence_and_reuses_manifest(self):
        job_id = "b" * 32
        path = pipeline.folder(job_id)
        path.mkdir(exist_ok=True)
        (path / "original").write_bytes(b"mock")
        job = {"id": job_id, "title": "Long audio"}
        pipeline.save_job(job)
        def fake_ffmpeg(*arguments):
            Path(arguments[-1]).write_bytes(b"prepared")
        with patch.object(pipeline, "duration", return_value=1250), patch.object(pipeline.subprocess, "run", return_value=SimpleNamespace(stderr="silence_end: 620.2\nsilence_end: 1190.5")), patch.object(pipeline, "run_ffmpeg", side_effect=fake_ffmpeg):
            parts = pipeline.split_audio(job)
        self.assertEqual([part["start"] for part in parts], [0.0, 620.2, 1190.5])
        with patch.object(pipeline, "run_ffmpeg", side_effect=AssertionError("should use existing parts")):
            self.assertEqual(pipeline.split_audio(pipeline.read_job(job_id)), parts)

    def test_multilingual_request_is_valid_multipart(self):
        job_id = "c" * 32
        path = pipeline.folder(job_id)
        path.mkdir(exist_ok=True)
        (path / "part-0000.m4a").write_bytes(b"mock")
        def inspect_request(_path, **kwargs):
            request = httpx.Client().build_request("POST", "https://example.test", data=kwargs["data"], files=kwargs["files"])
            body = request.read()
            self.assertIn(b'name="languages[]"\r\n\r\ncmn', body)
            self.assertIn(b'name="languages[]"\r\n\r\nyue', body)
            self.assertIn(b'name="file"', body)
            return {"text": "Hello, 你好。"}
        with patch.object(pipeline, "_post", side_effect=inspect_request):
            result = pipeline.transcribe_part({"id": job_id, "mode": "personal", "terms": ""}, {"file": "part-0000.m4a", "index": 0, "start": 0, "seconds": 10})
        self.assertEqual(result["segments"][0]["text"], "Hello, 你好。")

    def test_reports_disable_openai_application_state(self):
        def inspect_request(_path, **kwargs):
            self.assertIs(kwargs["payload"]["store"], False)
            return {"status": "completed", "output": [{"content": [{"type": "output_text", "text": '{"notes":"ok"}'}]}]}
        schema = {"type": "object", "properties": {"notes": {"type": "string"}}, "required": ["notes"], "additionalProperties": False}
        with patch.object(pipeline, "_post", side_effect=inspect_request):
            self.assertEqual(pipeline._json_response("sample", schema), {"notes": "ok"})

    def test_reviewed_speaker_names_appear_in_report(self):
        job = {"report": {"brief_en": "Part 1 · speaker_0 agreed."},
               "speaker_names": {"Part 1 · speaker_0": "Jayson"}}
        self.assertEqual(pipeline.named_report(job)["brief_en"], "Jayson agreed.")
        self.assertEqual(job["report"]["brief_en"], "Part 1 · speaker_0 agreed.")

    def test_translation_resumes_without_repeating_saved_sections(self):
        job_id = "d" * 32
        pipeline.folder(job_id).mkdir(exist_ok=True)
        # More than one request, containing a number, negation and both languages.
        originals = [("Budget is RM50,000; not approved. 还没批准。 " * 65) for _ in range(6)]
        job = {"id": job_id, "title": "Translation recovery", "transcribed": [{"segments": [
            {"start": i*60, "end": (i+1)*60, "speaker": "A", "text": text} for i, text in enumerate(originals)]}]}
        pipeline.save_job(job)
        groups = pipeline.translation_groups(pipeline.transcript_rows(job))
        calls = []
        def response(prompt, schema):
            import json
            rows = json.loads(prompt.split("\n\n", 1)[1])
            calls.append([row["id"] for row in rows])
            if len(calls) == 2:
                raise RuntimeError("network interrupted")
            return {"rows": [{"id": row["id"], "en": row["original"], "zh": row["original"]} for row in rows]}
        with patch.object(pipeline, "_json_response", side_effect=response):
            with self.assertRaisesRegex(RuntimeError, "interrupted"):
                pipeline.translate_transcript(job_id)
        self.assertEqual(len(pipeline.read_job(job_id)["translation_batches"]), 1)
        with patch.object(pipeline, "_json_response", side_effect=response):
            completed = pipeline.translate_transcript(job_id)
        self.assertEqual(calls.count([row["id"] for row in groups[0]]), 1)
        self.assertEqual(len(completed["translation_batches"]), len(groups))
        self.assertEqual(completed["transcribed"], job["transcribed"])
        self.assertEqual("".join(row["original"] for row in pipeline.transcript_rows(completed)), "".join(originals))
        self.assertTrue(all(row["en"] and row["zh"] for row in pipeline.transcript_rows(completed)))

    def test_incomplete_translation_does_not_advance_checkpoint(self):
        job_id = "e" * 32
        pipeline.folder(job_id).mkdir(exist_ok=True)
        job = {"id": job_id, "title": "Missing translation", "transcribed": [{"segments": [
            {"start": 0, "end": 1, "speaker": "A", "text": "Not agreed. 未达成共识。"}]}]}
        pipeline.save_job(job)
        for answer in ({"rows": []}, {"rows": [{"id": "wrong-id", "en": "test", "zh": "测试"}]},
                       {"rows": [{"id": "0-0-0", "en": "", "zh": "测试"}]}):
            with patch.object(pipeline, "_json_response", return_value=answer):
                with self.assertRaisesRegex(ValueError, "incomplete"):
                    pipeline.translate_transcript(job_id)
            self.assertFalse(pipeline.read_job(job_id).get("translation_batches"))

    def test_bilingual_exports_and_nested_speaker_names(self):
        from docx import Document
        job_id = "f" * 32
        pipeline.folder(job_id).mkdir(exist_ok=True)
        job = {"id": job_id, "title": "Family discussion", "filename": "family.m4a", "size": 8, "offset": 8,
               "status": "ready", "created": 1, "speaker_names": {"A": "Jayson"},
               "transcribed": [{"segments": [{"start": 10, "end": 20, "speaker": "A", "text": "We have not agreed. 还没决定。"}]}],
               "translation_batches": [[{"id": "0-0-0", "en": "We have not agreed. We have not decided.", "zh": "我们还未达成共识，也还没决定。"}]],
               "report": {"brief_en": "Not decided.", "brief_zh": "还没决定。", "comprehensive_en": "Discussion.", "comprehensive_zh": "讨论。",
                          "highlights_en": {"actions": [{"task": "Review options", "owner": "A", "deadline": ""}], "decisions": [], "questions": []}}}
        pipeline.save_job(job)
        pipeline.write_exports(job)
        self.assertEqual(pipeline.named_report(job)["highlights_en"]["actions"][0]["owner"], "Jayson")
        exported = (pipeline.folder(job_id) / "bilingual-transcript.txt").read_text()
        self.assertIn("[00:00:10–00:00:20] Jayson", exported)
        self.assertIn("Original: We have not agreed. 还没决定。", exported)
        with zipfile.ZipFile(pipeline.folder(job_id) / "memo-downloads.zip") as bundle:
            self.assertIn("bilingual-transcript.docx", bundle.namelist())
        doc = Document(pipeline.folder(job_id) / "brief.docx")
        self.assertEqual(doc.tables[0].rows[1].cells[1].text, "Jayson")
        self.assertEqual(doc.tables[0].rows[1].cells[2].text, "Not stated")
        response = self.client.get(f"/api/jobs/{job_id}/downloads/bilingual-transcript.docx")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Family", response.headers["content-disposition"])
        self.assertEqual(response.headers["cache-control"], "no-store")

    def test_private_endpoints_assets_and_legacy_translation_download(self):
        anonymous = TestClient(service.app, base_url="https://testserver")
        self.assertEqual(anonymous.get("/api/jobs").status_code, 401)
        for name in ("studio.css", "studio.js", "icon.svg", "manifest.webmanifest"):
            self.assertEqual(anonymous.get("/assets/" + name).status_code, 200)
        self.assertEqual(anonymous.get("/assets/app.py").status_code, 404)
        created = self.client.post("/api/jobs", json={"title": "Legacy", "filename": "x.m4a", "size": 1, "mode": "meeting"}, headers=self.headers).json()
        pipeline.update(created["id"], status="ready")
        self.assertEqual(self.client.get(f"/api/jobs/{created['id']}/downloads/bilingual-transcript.txt").status_code, 404)
        self.assertEqual(self.client.post("/api/logout", headers=self.headers).status_code, 200)
        self.assertEqual(self.client.get("/api/jobs").status_code, 401)
        anonymous.close()

    def test_speaker_rename_does_not_replace_words_or_cascade(self):
        self.assertEqual(pipeline._replace_speakers("A agreed. Approved by B.", {"A": "B", "B": "Mei"}), "B agreed. Approved by Mei.")
        self.assertEqual(pipeline._replace_speakers("Part 1 · speaker_01", {"Part 1 · speaker_0": "Jayson"}), "Part 1 · speaker_01")

    def test_process_can_finish_after_translation_interruption(self):
        job_id = "1234" * 8
        path = pipeline.folder(job_id)
        path.mkdir(exist_ok=True)
        job = {"id": job_id, "title": "Resume", "include_translation": True, "transcribed": [{"segments": [
            {"start": 0, "end": 10, "speaker": "A", "text": "No booking. 不要预订。"}]}]}
        pipeline.save_job(job)
        parts = [{"index": 0, "start": 0, "seconds": 10}]
        report = {key: "Report" for key in ("title_en", "title_zh", "brief_en", "brief_zh", "comprehensive_en", "comprehensive_zh")}
        with patch.object(pipeline, "split_audio", return_value=parts), patch.object(pipeline, "transcribe_part", side_effect=AssertionError("Already transcribed")), patch.object(pipeline, "_json_response", side_effect=RuntimeError("Translation interrupted")):
            pipeline.process(job_id)
        self.assertEqual(pipeline.read_job(job_id)["status"], "failed")
        with patch.object(pipeline, "split_audio", return_value=parts), patch.object(pipeline, "transcribe_part", side_effect=AssertionError("Already transcribed")), patch.object(pipeline, "_json_response", return_value={"rows": [{"id": "0-0-0", "en": "No booking. Do not book.", "zh": "不要预订。"}]}), patch.object(pipeline, "make_report", return_value=report):
            pipeline.process(job_id)
        finished = pipeline.read_job(job_id)
        self.assertEqual(finished["status"], "ready")
        self.assertEqual(finished["transcribed"], job["transcribed"])
        self.assertTrue((path / "bilingual-transcript.docx").is_file())


if __name__ == "__main__":
    unittest.main()
